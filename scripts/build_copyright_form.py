"""软件著作权登记信息表生成

用法: python scripts/build_copyright_form.py
产出: docs/软著-登记信息表.docx
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "软著-登记信息表.docx"

FIELDS = [
    ("软件全称", "义乌小商品出海智能体平台"),
    ("软件简称", "义乌出海智能体"),
    ("版本号", "V1.0"),
    ("开发完成日期", date.today().isoformat()),
    ("首次发表日期", "未发表"),
    ("著作权人", "（按申请主体填写）"),
    ("开发方式", "独立开发"),
    ("软件用途", "面向跨境电商的一站式AI智能服务平台：为义乌小商品商户提供市场洞察、智能选品、供应链匹配、跨境内容生成、合规查询、智能客服、政策复制七大AI Agent全链路服务，助力义乌经验向全国39个市场采购贸易试点城市复制推广。"),
    ("技术特点", "基于 FastAPI + LangGraph 的多Agent编排架构（StateGraph条件路由）；LLM 大模型增强（主备双模自动降级、推理模型空内容自愈）；React 18 前端；SQLite 持久化；103 个自动化单元测试全绿；Docker 容器化部署于魔搭创空间并具备 CI/CD 流水线与部署后 AI 生效性冒烟测试。"),
    ("编程语言", "Python 3.12 / TypeScript（React 18）"),
    ("源程序量", "约 5,535 行（Python 后端+测试），另有 TypeScript 前端约 8,000 行"),
    ("运行环境", "服务端：Linux/Windows，Python 3.12+，2vCPU/16GB 内存容器；客户端：现代浏览器（Chrome/Edge）"),
    ("软件功能模块", "①市场洞察Agent ②智能选品Agent ③供应链匹配Agent ④跨境内容生成Agent ⑤合规助手Agent ⑥智能客服Agent ⑦政策复制Agent ⑧用户认证与用量统计 ⑨LLM服务层（主备降级+限额）"),
]


def build():
    doc = Document()
    doc.add_heading("计算机软件著作权登记信息表", level=1, )

    table = doc.add_table(rows=len(FIELDS), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(FIELDS):
        c0, c1 = table.rows[i].cells
        c0.width, c1.width = Cm(4.5), Cm(12)
        r0 = c0.paragraphs[0].add_run(k)
        r0.bold = True
        r0.font.size = Pt(10.5)
        r1 = c1.paragraphs[0].add_run(v)
        r1.font.size = Pt(10.5)

    doc.add_paragraph()
    doc.add_paragraph("备注：源代码文档（每页50行，前30页+后30页）见《软著-源代码文档.docx》；程序说明书见《用户使用手册》。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"登记信息表: {OUT}")


if __name__ == "__main__":
    build()
